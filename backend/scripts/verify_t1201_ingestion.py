"""T1201 — isolated ingestion-pipeline E2E verification.

Run from the repository root::

    python backend/scripts/verify_t1201_ingestion.py

The child-process matrix exercises the real FastAPI upload endpoint, parser
dispatch, text cleaning, chunking, embedding-service boundary, ChromaDB
persistence, file listing, and keyword-index invalidation. Fixtures are built
at runtime for every supported extension, including native, scanned, mixed,
partially failing, and fully failing PDFs.

Isolation and declared substitutions
------------------------------------
``UPLOAD_DIR`` and ``CHROMA_PERSIST_DIR`` point to a fresh system temp tree,
which the parent removes after the child exits and releases ChromaDB's Windows
file handles. The repository's real uploads and Chroma data are never touched.

The repository-local BGE model is verified separately by
``verify_bge_model.py``. This broad deterministic matrix substitutes the BGE
inference edge and the unapproved DashScope edge so it can cover all ingestion
branches without provider calls or repeatedly running heavyweight inference:

* ``SentenceTransformer`` construction returns a deterministic 512-dimension,
  L2-normalized model double while the real lazy singleton and
  ``encode_chunks`` implementation still run.
* ``MultiModalConversation.call`` returns deterministic responses (or 5xx)
  while the real PDF page detection, rendering, Base64 request assembly,
  retry/warning logic, cleaning, chunking, API mapping, and rollback run.

Exit code 0 means every required check passed; 1 means at least one failed.

Explicitly authorized live verification::

    python backend/scripts/verify_t1201_ingestion.py --live

Live mode reads backend/.env without printing it, runs the real local BGE
model and real DashScope SDK, and records only allowlisted response evidence.
Failure cases use a real SDK request with a 1-microsecond network timeout;
no response/exception is fabricated and the application's real backoff runs.
These are controlled transport failures, not observed provider 5xx outages.
"""

from __future__ import annotations

import shutil
import subprocess
import sys
import tempfile
import time
from pathlib import Path

_BACKEND = Path(__file__).resolve().parents[1]


def run_probe(probe_root: Path, *, live: bool = False) -> int:
    """Run the verification matrix in the isolated child process."""
    import ast
    import hashlib
    import math
    import os
    import types
    import uuid
    import logging
    from datetime import datetime, timezone
    from contextlib import contextmanager, nullcontext
    from types import SimpleNamespace
    from unittest.mock import patch

    os.environ["ANONYMIZED_TELEMETRY"] = "False"
    os.environ["UPLOAD_DIR"] = str(probe_root / "uploads")
    os.environ["CHROMA_PERSIST_DIR"] = str(probe_root / "chroma")
    if live:
        os.chdir(_BACKEND)
        os.environ["EMBED_MODEL"] = str(_BACKEND / "models" / "bge-small-zh-v1.5")
        os.environ["HF_HUB_OFFLINE"] = "1"
        os.environ["TRANSFORMERS_OFFLINE"] = "1"
        logging.disable(logging.CRITICAL)
    else:
        os.environ["DASHSCOPE_API_KEY"] = f"test-double-{probe_root.name}"

    sys.path.insert(0, str(_BACKEND))

    import dashscope
    import fitz
    from docx import Document
    from fastapi.testclient import TestClient
    from openpyxl import Workbook

    import app.services.embedding as embedding_mod
    import app.services.ingest as ingest_mod
    from app.core.config import settings
    from app.core.vector_store import ChromaVectorStore
    from app.main import app
    from app.services.qa import KeywordRetriever

    results: list[tuple[str, bool, str]] = []
    supported_successes: set[str] = set()
    print("EVIDENCE_MODE:", "REAL BGE + REAL DashScope; controlled network timeouts" if live else "SUBSTITUTED BGE + OCR")
    print("STARTED_UTC:", datetime.now(timezone.utc).isoformat())
    if live and not settings.get_dashscope_key():
        print("BLOCKED: DASHSCOPE_API_KEY is not configured")
        return 1
    if live:
        from dotenv import dotenv_values
        if settings.get_dashscope_key() != dotenv_values(_BACKEND / ".env").get("DASHSCOPE_API_KEY"):
            print("BLOCKED: effective DashScope credential differs from backend/.env")
            return 1

    def check(label: str, condition: bool, detail: str = "") -> None:
        ok = bool(condition)
        results.append((label, ok, detail))
        print(f"    [{'PASS' if ok else 'FAIL'}] {label}" + (f" — {detail}" if detail else ""))

    def body_of(response) -> dict:
        try:
            value = response.json()
        except Exception:
            return {}
        return value if isinstance(value, dict) else {}

    def code_of(response) -> str | None:
        return body_of(response).get("error", {}).get("code")

    def is_uuid(value: object) -> bool:
        try:
            uuid.UUID(str(value))
            return True
        except (TypeError, ValueError, AttributeError):
            return False

    class Matrix:
        def __init__(self, rows: list[list[float]]) -> None:
            self.rows = rows

        def tolist(self) -> list[list[float]]:
            return self.rows

    class DeterministicSentenceTransformer:
        load_count = 0
        loaded_paths: list[str] = []

        def __init__(self, model_path: str) -> None:
            type(self).load_count += 1
            type(self).loaded_paths.append(model_path)
            self.encode_calls: list[tuple[list[str], bool]] = []

        def encode(self, texts, normalize_embeddings=False):
            values = list(texts)
            self.encode_calls.append((values, normalize_embeddings))
            rows = []
            for text in values:
                digest = hashlib.sha256(text.encode("utf-8")).digest()
                index = (
                    int.from_bytes(digest[:2], "big")
                    % embedding_mod.EMBEDDING_DIMENSION
                )
                vector = [0.0] * embedding_mod.EMBEDDING_DIMENSION
                vector[index] = 1.0
                rows.append(vector)
            return Matrix(rows)

    fake_sentence_transformers = types.ModuleType("sentence_transformers")
    fake_sentence_transformers.SentenceTransformer = DeterministicSentenceTransformer
    embedding_mod._model = None
    with (nullcontext() if live else patch.dict(sys.modules, {"sentence_transformers": fake_sentence_transformers})):
        first_model = embedding_mod.get_model()
        second_model = embedding_mod.get_model()
        contract_vectors = embedding_mod.encode_chunks(["one", "two", "three"])

    check(
        "AC-F007-02 lazy model is constructed once and reused",
        first_model is second_model and (live or DeterministicSentenceTransformer.load_count == 1),
        "same model object reused" if live else f"loads={DeterministicSentenceTransformer.load_count}",
    )
    check(
        "AC-F007-02 configured model path reaches the loader",
        Path(settings.EMBED_MODEL).is_dir() if live else DeterministicSentenceTransformer.loaded_paths == [settings.EMBED_MODEL],
        f"path={settings.EMBED_MODEL}",
    )
    check(
        "AC-F007-01 embedding count, width, and L2 normalization",
        len(contract_vectors) == 3
        and all(
            len(vector) == embedding_mod.EMBEDDING_DIMENSION
            for vector in contract_vectors
        )
        and all(math.isclose(math.sqrt(sum(value * value for value in vector)), 1.0, abs_tol=1e-6) for vector in contract_vectors)
        and (live or first_model.encode_calls[-1][1] is True),
    )

    if live:
        norms = [round(math.sqrt(sum(v*v for v in row)), 7) for row in contract_vectors]
        print(f"BGE_MODEL class={type(first_model).__module__}.{type(first_model).__name__} shape={len(contract_vectors)}x{len(contract_vectors[0])} norms={norms}")

    repo_root = _BACKEND.parent.resolve()
    upload_root = Path(settings.UPLOAD_DIR).resolve()
    chroma_root = Path(settings.CHROMA_PERSIST_DIR).resolve()
    isolated_storage = (
        upload_root != repo_root
        and chroma_root != repo_root
        and repo_root not in upload_root.parents
        and repo_root not in chroma_root.parents
    )
    check(
        "probe storage is outside the repository",
        isolated_storage,
        f"uploads={upload_root}; chroma={chroma_root}",
    )
    if not isolated_storage:
        raise RuntimeError("refusing to run T1201 verification inside the repository")

    store = ChromaVectorStore()
    client = TestClient(app, raise_server_exceptions=False)
    fixture_root = probe_root / "fixtures"
    fixture_root.mkdir(parents=True, exist_ok=True)

    def create_kb(name: str) -> None:
        response = client.post("/api/collections", json={"name": name})
        check(f"create isolated knowledge base {name}", response.status_code == 201)

    def upload(kb: str, name: str, content: bytes, content_type: str = "application/octet-stream"):
        return client.post(
            "/api/upload",
            files={"file": (name, content, content_type)},
            data={"collection_name": kb},
        )

    def files_in(kb: str) -> list[dict]:
        response = client.get("/api/files", params={"collection_name": kb})
        check(f"file-list API responds for {kb}", response.status_code == 200)
        return body_of(response).get("files", [])

    def chunks_for(kb: str, file_id: str):
        return store.get_chunks_by_file(kb, file_id)

    def stored_text(kb: str, file_id: str) -> str:
        return "\n".join(chunk.content for chunk in chunks_for(kb, file_id))

    def record_success(response, extension: str, label: str) -> dict:
        body = body_of(response)
        ok = (
            response.status_code == 200
            and body.get("status") in {"SUCCESS", "SUCCESS_WITH_WARNINGS"}
            and body.get("chunks", 0) > 0
            and is_uuid(body.get("file_id"))
        )
        check(label, ok, f"HTTP {response.status_code} {body.get('status')}")
        if ok:
            supported_successes.add(extension)
        return body

    def docx_fixture() -> bytes:
        path = fixture_root / "table.docx"
        document = Document()
        document.add_paragraph("Document introduction")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Name"
        table.cell(0, 1).text = "Value"
        table.cell(1, 0).text = "Alpha"
        table.cell(1, 1).text = "42"
        document.save(path)
        return path.read_bytes()

    def workbook_fixture(extension: str, *, multi_sheet: bool) -> bytes:
        path = fixture_root / f"workbook{extension}"
        workbook = Workbook()
        first = workbook.active
        first.title = "Sheet1"
        first.append(["Alpha", 1])
        if multi_sheet:
            workbook.create_sheet("Sheet2")
            third = workbook.create_sheet("Sheet3")
            third.append(["Gamma", 3])
        if extension in {".xltx", ".xltm"}:
            workbook.template = True
        workbook.save(path)
        return path.read_bytes()

    def raster_page_image(text: str) -> bytes:
        source = fitz.open()
        page = source.new_page(width=595, height=842)
        page.insert_text((72, 100), text, fontsize=18)
        image = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5)).tobytes("png")
        source.close()
        return image

    def pdf_fixture(name: str, pages: list[tuple[str, str]]) -> bytes:
        path = fixture_root / name
        document = fitz.open()
        for kind, text in pages:
            page = document.new_page(width=595, height=842)
            if kind == "native":
                page.insert_text((72, 100), text, fontsize=14)
            else:
                page.insert_image(page.rect, stream=raster_page_image(text))
        document.save(path)
        document.close()
        content = path.read_bytes()
        if live:
            print(f"FIXTURE name={name} sha256={hashlib.sha256(content).hexdigest()} pages={pages!r}")
        reopened = fitz.open(stream=content, filetype="pdf")
        for index, (kind, _) in enumerate(pages):
            if kind == "scan":
                check(
                    f"{name} page {index + 1} has no native text",
                    not reopened[index].get_text().strip(),
                )
        reopened.close()
        return content

    def ocr_success(text: str):
        return SimpleNamespace(
            status_code=200,
            output=SimpleNamespace(
                choices=[
                    SimpleNamespace(
                        message=SimpleNamespace(content=[{"text": text}])
                    )
                ]
            ),
        )

    def ocr_failure():
        return SimpleNamespace(status_code=500)

    @contextmanager
    def observe_ocr(*, texts=None, fail=False):
        if not live:
            kwargs = {"return_value": ocr_failure()} if fail else {"side_effect": [ocr_success(text) for text in (texts or [])]}
            with patch.object(dashscope.MultiModalConversation, "call", **kwargs) as observed:
                yield observed
            return
        # This observer delegates every attempt to the real SDK. It never
        # changes credentials, endpoint, model, messages, or returned output.
        descriptor = dashscope.MultiModalConversation.__dict__["call"]
        original = dashscope.MultiModalConversation.call
        observed = SimpleNamespace(call_count=0, call_args_list=[], failures=[])

        def call(**kwargs):
            observed.call_count += 1
            observed.call_args_list.append(SimpleNamespace(kwargs=kwargs))
            started = time.monotonic()
            try:
                response = original(**kwargs, request_timeout=0.000001 if fail else 60)
            except Exception as exc:
                # Exception messages / request headers must never enter logs.
                observed.failures.append(type(exc).__name__)
                print(f"LIVE_OCR attempt={observed.call_count} controlled_timeout={fail} exception={type(exc).__name__} elapsed={time.monotonic()-started:.3f}s")
                cause = exc.__cause__ or exc.__context__
                if cause is not None:
                    print(f"LIVE_TRANSPORT_CAUSE type={type(cause).__name__} errno={getattr(cause, 'errno', None)}")
                raise
            print(f"LIVE_OCR attempt={observed.call_count} controlled_timeout={fail} status={response.status_code} request_id={response.request_id} elapsed={time.monotonic()-started:.3f}s")
            return response

        dashscope.MultiModalConversation.call = staticmethod(call)
        try:
            yield observed
        finally:
            dashscope.MultiModalConversation.call = descriptor
            if fail:
                check("controlled failure is a real SDK transport timeout", bool(observed.failures) and all("Timeout" in name for name in observed.failures), f"exceptions={observed.failures}")

    @contextmanager
    def observe_sleep():
        if not live:
            with patch.object(ingest_mod.time, "sleep", return_value=None) as observed:
                yield observed
            return
        original = ingest_mod.time.sleep
        observed = SimpleNamespace(call_count=0)

        def sleep(seconds):
            observed.call_count += 1
            print(f"LIVE_BACKOFF seconds={seconds}")
            original(seconds)

        ingest_mod.time.sleep = sleep
        try:
            yield observed
        finally:
            ingest_mod.time.sleep = original

    print("\n--- Text formats, cleaning, and plain-text dispatch")
    create_kb("t1201-text")
    utf8_response = upload(
        "t1201-text",
        "utf8.txt",
        "  第一行  \n\n  第二行  ".encode("utf-8"),
        "text/plain",
    )
    utf8_body = record_success(utf8_response, ".txt", "AC-F003-01 UTF-8 TXT upload")
    check(
        "AC-F005-01 cleaning removes blank lines and edge whitespace",
        stored_text("t1201-text", utf8_body.get("file_id", "")) == "第一行\n第二行",
    )

    gbk_text = "中文 GBK 编码内容"
    gbk_response = upload("t1201-text", "gbk.txt", gbk_text.encode("gbk"), "text/plain")
    gbk_body = record_success(gbk_response, ".txt", "AC-F003-01/02 GBK fallback upload")
    check(
        "AC-F003-02 GBK text remains readable after the full pipeline",
        gbk_text in stored_text("t1201-text", gbk_body.get("file_id", "")),
    )
    utf16_body = record_success(upload("t1201-text", "utf16.txt", gbk_text.encode("utf-16")), ".txt", "F003 UTF-16 fallback")
    check("UTF-16 decoded text is preserved", stored_text("t1201-text", utf16_body.get("file_id", "")) == gbk_text)

    plain_formats = {
        ".csv": b"name,value\nalpha,1",
        ".json": b'{"name":"alpha","value":1}',
        ".log": b"2026-09-07 ingestion started",
    }
    for extension, content in plain_formats.items():
        response = upload("t1201-text", f"plain{extension}", content, "text/plain")
        record_success(response, extension, f"supported plain-text format {extension}")

    blank_response = upload("t1201-text", "whitespace.txt", b" \n\t\n ", "text/plain")
    check(
        "AC-F005-02 cleaned-empty upload is rejected",
        blank_response.status_code == 422 and code_of(blank_response) == "FILE_PARSE_ERROR",
        f"HTTP {blank_response.status_code} {code_of(blank_response)}",
    )
    check(
        "AC-F005-02 cleaned-empty upload leaves no raw file",
        not (upload_root / "t1201-text" / "whitespace.txt").exists(),
    )

    print("\n--- Markdown and recursive chunking")
    create_kb("t1201-chunk")
    markdown = "## 章节A\n正文 A\n\n### 小节A1\n内容 A1"
    md_response = upload("t1201-chunk", "headings.md", markdown.encode("utf-8"), "text/markdown")
    md_body = record_success(md_response, ".md", "AC-F006-01 Markdown upload")
    md_chunks = chunks_for("t1201-chunk", md_body.get("file_id", ""))
    check(
        "AC-F006-01 nested header path prefixes a chunk",
        any(chunk.content.startswith("章节A > 小节A1\n\n内容 A1") for chunk in md_chunks),
        f"chunks={[chunk.content for chunk in md_chunks]}",
    )

    short_text = "短" * 300
    short_response = upload("t1201-chunk", "short.txt", short_text.encode("utf-8"), "text/plain")
    short_body = record_success(short_response, ".txt", "AC-F006-03 short text upload")
    short_chunks = chunks_for("t1201-chunk", short_body.get("file_id", ""))
    check(
        "AC-F006-03 300-character text remains one chunk",
        len(short_chunks) == 1 and short_chunks[0].content == short_text,
    )

    long_text = "".join(chr(0x4E00 + (index % 1000)) for index in range(2000))
    long_response = upload("t1201-chunk", "long.txt", long_text.encode("utf-8"), "text/plain")
    long_body = record_success(long_response, ".txt", "AC-F006-02 long text upload")
    long_chunks = chunks_for("t1201-chunk", long_body.get("file_id", ""))
    check(
        "AC-F006-02 recursive chunks stay within 800 characters",
        len(long_chunks) > 1 and all(len(chunk.content) <= settings.MAX_CHUNK_SIZE for chunk in long_chunks),
        f"sizes={[len(chunk.content) for chunk in long_chunks]}",
    )
    check(
        "AC-F006-02 adjacent recursive chunks overlap by 120 characters",
        len(long_chunks) > 1
        and long_chunks[0].content[-settings.CHUNK_OVERLAP :]
        == long_chunks[1].content[: settings.CHUNK_OVERLAP],
    )

    print("\n--- DOCX and all Excel variants")
    create_kb("t1201-office")
    docx_response = upload(
        "t1201-office",
        "table.docx",
        docx_fixture(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    docx_body = record_success(docx_response, ".docx", "AC-F003-03/04 DOCX table upload")
    docx_text = stored_text("t1201-office", docx_body.get("file_id", ""))
    check(
        "AC-F003-03/04 DOCX paragraph and table cells are extracted",
        "Document introduction" in docx_text
        and "Name Value" in docx_text
        and "Alpha 42" in docx_text,
        f"content={docx_text!r}",
    )

    excel_extensions = [".xlsx", ".xlsm", ".xltx", ".xltm"]
    for extension in excel_extensions:
        response = upload(
            "t1201-office",
            f"book{extension}",
            workbook_fixture(extension, multi_sheet=extension == ".xlsx"),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        body = record_success(response, extension, f"supported Excel format {extension}")
        text = stored_text("t1201-office", body.get("file_id", ""))
        check(f"{extension} cell content extracted", "Alpha 1" in text)
        if extension == ".xlsx":
            check(
                "AC-F003-05 XLSX skips empty sheet and reads later sheet",
                "Alpha 1" in text and "Gamma 3" in text,
                f"content={text!r}",
            )

    print("\n--- Native, scanned, mixed, warning, and FAILED PDFs")
    create_kb("t1201-pdf")
    native_pdf = pdf_fixture("native.pdf", [("native", "NATIVE_PDF_TEXT"), ("native", "NATIVE_SECOND_PAGE")])
    with observe_ocr() as ocr_call:
        native_response = upload("t1201-pdf", "native.pdf", native_pdf, "application/pdf")
    native_body = record_success(native_response, ".pdf", "AC-F002-01 native PDF upload")
    check("native PDF does not call OCR", ocr_call.call_count == 0)
    check(
        "AC-F002-01 raw PDF, chunks, and file-list record persist",
        (upload_root / "t1201-pdf" / "native.pdf").exists()
        and native_body.get("chunks", 0) == len(chunks_for("t1201-pdf", native_body.get("file_id", "")))
        and any(item.get("file_id") == native_body.get("file_id") for item in files_in("t1201-pdf")),
    )
    check(
        "native PDF text survives parsing and storage",
        "NATIVE_PDF_TEXT\nNATIVE_SECOND_PAGE" == stored_text("t1201-pdf", native_body.get("file_id", "")),
    )

    scanned_pdf = pdf_fixture(
        "scanned.pdf",
        [("scan", "SCAN IMAGE ONE"), ("scan", "SCAN IMAGE TWO")],
    )
    with observe_ocr(texts=["OCR_PAGE_ONE", "OCR_PAGE_TWO"]) as ocr_call:
        scanned_response = upload("t1201-pdf", "scanned.pdf", scanned_pdf, "application/pdf")
    scanned_body = record_success(scanned_response, ".pdf", "AC-F004-01 scanned PDF OCR success")
    scanned_text = stored_text("t1201-pdf", scanned_body.get("file_id", ""))
    if live:
        print(f"OCR_STORED scanned={scanned_text!r}")
    check(
        "AC-F004-01 scanned pages preserve OCR order with no warnings",
        scanned_response.status_code == 200
        and scanned_body.get("status") == "SUCCESS"
        and scanned_body.get("warnings") == []
        and ("SCAN IMAGE ONE" if live else "OCR_PAGE_ONE") in scanned_text
        and ("SCAN IMAGE TWO" if live else "OCR_PAGE_TWO") in scanned_text
        and scanned_text.index("SCAN IMAGE ONE" if live else "OCR_PAGE_ONE") < scanned_text.index("SCAN IMAGE TWO" if live else "OCR_PAGE_TWO")
        and ocr_call.call_count == 2,
    )
    first_ocr_kwargs = ocr_call.call_args_list[0].kwargs if ocr_call.call_args_list else {}
    first_ocr_content = first_ocr_kwargs.get("messages", [{}])[0].get("content", [])
    check(
        "F004 real OCR adapter sends qwen-vl-plus, data URI, and frozen prompt",
        first_ocr_kwargs.get("model") == "qwen-vl-plus"
        and first_ocr_content[0].get("image", "").startswith("data:image/jpeg;base64,")
        and first_ocr_content[1].get("text") == "请提取图片中的所有文字，保持格式",
    )

    mixed_pdf = pdf_fixture(
        "mixed.pdf",
        [("native", "MIXED_NATIVE_FIRST"), ("scan", "MIXED SCAN SECOND")],
    )
    with observe_ocr(texts=["MIXED_OCR_SECOND"]) as ocr_call:
        mixed_response = upload("t1201-pdf", "mixed.pdf", mixed_pdf, "application/pdf")
    mixed_body = record_success(mixed_response, ".pdf", "AC-F003-02 / AC-F004-02 mixed PDF")
    mixed_text = stored_text("t1201-pdf", mixed_body.get("file_id", ""))
    if live:
        print(f"OCR_STORED mixed={mixed_text!r}")
    check(
        "AC-F003-02 / AC-F004-02 native and OCR pages are ordered and OCR is selective",
        mixed_body.get("status") == "SUCCESS"
        and "MIXED_NATIVE_FIRST" in mixed_text
        and ("MIXED SCAN SECOND" if live else "MIXED_OCR_SECOND") in mixed_text
        and mixed_text.index("MIXED_NATIVE_FIRST") < mixed_text.index("MIXED SCAN SECOND" if live else "MIXED_OCR_SECOND")
        and ocr_call.call_count == 1,
    )

    partial_pdf = pdf_fixture(
        "partial.pdf",
        [
            ("native", "PARTIAL_PAGE_1"),
            ("native", "PARTIAL_PAGE_2"),
            ("scan", "PARTIAL SCAN PAGE 3"),
            ("native", "PARTIAL_PAGE_4"),
            ("native", "PARTIAL_PAGE_5"),
        ],
    )
    with observe_ocr(fail=True) as ocr_call, observe_sleep() as sleep_call:
        partial_response = upload("t1201-pdf", "partial.pdf", partial_pdf, "application/pdf")
    partial_body = record_success(
        partial_response,
        ".pdf",
        "AC-F002-07 / AC-F004-03 partial OCR failure",
    )
    expected_partial_warning = [{"page_number": 3, "error_code": "OCR_PAGE_FAILED"}]
    partial_text = stored_text("t1201-pdf", partial_body.get("file_id", ""))
    if live:
        print(f"PARTIAL_RESULT status={partial_body.get('status')} chunks={partial_body.get('chunks')} warnings={partial_body.get('warnings')} stored={partial_text!r}")
    check(
        "AC-F002-07 / AC-F004-03 returns SUCCESS_WITH_WARNINGS and keeps good pages",
        partial_response.status_code == 200
        and partial_body.get("status") == "SUCCESS_WITH_WARNINGS"
        and partial_body.get("warnings") == expected_partial_warning
        and all(f"PARTIAL_PAGE_{page}" in partial_text for page in (1, 2, 4, 5))
        and ocr_call.call_count == 3
        and sleep_call.call_count == 2,
    )
    check(
        "AC-F004-05 failed OCR page is never silently omitted",
        partial_body.get("warnings") == expected_partial_warning,
    )
    partial_duplicate = upload("t1201-pdf", "PARTIAL.PDF", partial_pdf, "application/pdf")
    check(
        "AC-F002-10 SUCCESS_WITH_WARNINGS remains duplicate-protected",
        partial_duplicate.status_code == 409 and code_of(partial_duplicate) == "FILE_ALREADY_EXISTS",
        f"HTTP {partial_duplicate.status_code} {code_of(partial_duplicate)}",
    )

    all_failed_pdf = pdf_fixture(
        "all-failed.pdf",
        [("scan", "FAILED SCAN 1"), ("scan", "FAILED SCAN 2"), ("scan", "FAILED SCAN 3")],
    )
    failed_baseline_count = store.get_chunk_count("t1201-pdf")
    failed_baseline_chunks = store.list_chunks("t1201-pdf")
    KeywordRetriever(store).keyword_search("t1201-pdf", "rollbackabsencetoken", 5)
    baseline_index = {token: set(ids) for token, ids in KeywordRetriever._indexes["t1201-pdf"].items()}
    with observe_ocr(fail=True) as ocr_call, observe_sleep():
        failed_response = upload("t1201-pdf", "all-failed.pdf", all_failed_pdf, "application/pdf")
    failed_warnings = [
        {"page_number": page, "error_code": "OCR_PAGE_FAILED"}
        for page in (1, 2, 3)
    ]
    if live:
        print(f"FAILED_RESULT http={failed_response.status_code} code={code_of(failed_response)} warnings={body_of(failed_response).get('error', {}).get('details', {}).get('warnings')}")
    check(
        "AC-F002-08 / AC-F004-04 all OCR pages failing returns 422 with warnings",
        failed_response.status_code == 422
        and code_of(failed_response) == "FILE_PARSE_ERROR"
        and body_of(failed_response).get("error", {}).get("details", {}).get("warnings") == failed_warnings
        and ocr_call.call_count == 9,
        f"HTTP {failed_response.status_code} {code_of(failed_response)}",
    )
    current_files = files_in("t1201-pdf")
    failed_keyword_results = KeywordRetriever(store).keyword_search(
        "t1201-pdf", "rollbackabsencetoken", 5
    )
    check(
        "AC-F002-09 FAILED leaves no raw file, chunk/vector/metadata, or keyword entry",
        not (upload_root / "t1201-pdf" / "all-failed.pdf").exists()
        and store.get_chunk_count("t1201-pdf") == failed_baseline_count
        and sorted(store.list_chunks("t1201-pdf"), key=lambda chunk: chunk.chunk_id) == sorted(failed_baseline_chunks, key=lambda chunk: chunk.chunk_id)
        and KeywordRetriever._indexes["t1201-pdf"] == baseline_index
        and set(KeywordRetriever._chunks["t1201-pdf"]) == {chunk.chunk_id for chunk in failed_baseline_chunks}
        and not any(item.get("file_name") == "all-failed.pdf" for item in current_files)
        and failed_keyword_results == [],
        f"chunks_before={failed_baseline_count}; chunks_after={store.get_chunk_count('t1201-pdf')}; raw_exists={(upload_root / 't1201-pdf' / 'all-failed.pdf').exists()}; keyword_query_matches={len(failed_keyword_results)}",
    )
    with observe_ocr(texts=["RECOVERED_1", "RECOVERED_2", "RECOVERED_3"]):
        retry_response = upload("t1201-pdf", "all-failed.pdf", all_failed_pdf, "application/pdf")
    check(
        "AC-F002-09 same-name retry succeeds after FAILED rollback",
        retry_response.status_code == 200 and body_of(retry_response).get("status") == "SUCCESS",
        f"HTTP {retry_response.status_code} {body_of(retry_response).get('status')}",
    )
    if live:
        recovered = stored_text("t1201-pdf", body_of(retry_response).get("file_id", ""))
        print(f"OCR_STORED recovered={recovered!r}")
        check("same-name retry really OCRs all three scan pages", all(f"FAILED SCAN {page}" in recovered for page in (1, 2, 3)))

    print("\n--- Validation, duplicate scope, and keyword-index invalidation")
    create_kb("t1201-valid")
    validation_baseline = store.get_chunk_count("t1201-valid")
    rejected = [
        ("bad.exe", b"payload", 400, "UNSUPPORTED_FILE_TYPE"),
        ("empty.txt", b"", 400, "EMPTY_FILE"),
        ("oversized.txt", b"x" * (51 * 1024 * 1024), 413, "FILE_TOO_LARGE"),
        ("../escape.txt", b"payload", 400, "INVALID_FILE_NAME"),
        ("subdir/escape.txt", b"payload", 400, "INVALID_FILE_NAME"),
        ("broken.pdf", b"not a PDF", 422, "FILE_PARSE_ERROR"),
    ]
    for name, content, status, error_code in rejected:
        response = upload("t1201-valid", name, content)
        check(
            f"AC-F002 validation rejects {name}",
            response.status_code == status and code_of(response) == error_code,
            f"HTTP {response.status_code} {code_of(response)}",
        )
    check(
        "AC-F002-04/05/06 validation failures have no persistence side effects",
        store.get_chunk_count("t1201-valid") == validation_baseline
        and not any((upload_root / "t1201-valid" / name).exists() for name, *_ in rejected),
        f"remaining_files={[str(path.relative_to(upload_root)) for path in (upload_root / 't1201-valid').rglob('*') if path.is_file()]}",
    )

    create_kb("t1201-kb-a")
    create_kb("t1201-kb-b")
    same_a = upload("t1201-kb-a", "shared.txt", b"content in knowledge base A", "text/plain")
    same_b = upload("t1201-kb-b", "shared.txt", b"content in knowledge base B", "text/plain")
    check(
        "AC-F002-03 same filename succeeds independently in two knowledge bases",
        same_a.status_code == 200 and same_b.status_code == 200,
    )
    duplicate = upload("t1201-kb-a", "SHARED.TXT", b"duplicate", "text/plain")
    check(
        "AC-F002-02 case-variant duplicate in one knowledge base is rejected",
        duplicate.status_code == 409 and code_of(duplicate) == "FILE_ALREADY_EXISTS",
        f"HTTP {duplicate.status_code} {code_of(duplicate)}",
    )

    create_kb("t1201-index")
    retriever = KeywordRetriever(store)
    check("keyword index starts empty", retriever.keyword_search("t1201-index", "phoenixterm", 5) == [])
    indexed_response = upload(
        "t1201-index",
        "indexed.txt",
        "phoenixterm appears after the cached empty index".encode("utf-8"),
        "text/plain",
    )
    indexed_body = record_success(indexed_response, ".txt", "upload after empty keyword index cache")
    rebuilt_results = retriever.keyword_search("t1201-index", "phoenixterm", 5)
    check(
        "successful ingestion invalidates and rebuilds the keyword index",
        indexed_body.get("status") == "SUCCESS"
        and any(result.get("file_name") == "indexed.txt" for result in rebuilt_results),
    )

    print("\n--- Vector persistence contract through ingested chunks")
    create_kb("t1201-vector")
    sections = "\n\n".join(
        f"## Section {index}\nVector integration content {index} token-{index}"
        for index in range(10)
    )
    vector_response = upload(
        "t1201-vector", "ten.md", sections.encode("utf-8"), "text/markdown"
    )
    vector_body = record_success(vector_response, ".md", "ten-chunk vector fixture upload")
    vector_chunks = chunks_for("t1201-vector", vector_body.get("file_id", ""))
    check(
        "AC-F008-01 ingestion persisted ten chunks with immutable UUID identities",
        len(vector_chunks) == 10
        and all(is_uuid(chunk.chunk_id) and chunk.file_id == vector_body.get("file_id") for chunk in vector_chunks)
        and [chunk.chunk_index for chunk in vector_chunks] == list(range(10)),
        f"chunk_count={len(vector_chunks)}",
    )
    expected_metadata = {
        "chunk_id",
        "file_id",
        "file_name",
        "collection_name",
        "chunk_index",
        "source_file",
        "file_size",
        "upload_time",
        "ingestion_status",
    }
    check(
        "F008 persisted metadata has the exact nine required fields",
        bool(vector_chunks)
        and all(set(chunk.metadata) == expected_metadata for chunk in vector_chunks),
    )
    query_vector = embedding_mod.encode_chunks([vector_chunks[0].content])[0]
    search_results = store.search("t1201-vector", query_vector, top_k=3)
    check(
        "AC-F008-01 public search returns top-3 similarity results with identities",
        len(search_results) == 3
        and search_results[0].chunk_id == vector_chunks[0].chunk_id
        and all(0.0 <= item.similarity_score <= 1.0 for item in search_results)
        and all(item.file_id == vector_body.get("file_id") for item in search_results),
    )

    survivor_response = upload(
        "t1201-vector", "survivor.txt", b"unrelated survivor", "text/plain"
    )
    survivor_body = body_of(survivor_response)
    deleted_count = store.delete_by_file("t1201-vector", vector_body.get("file_id", ""))
    check(
        "AC-F008-02 delete_by_file removes one file's chunks and preserves another",
        deleted_count == 10
        and store.get_chunks_by_file("t1201-vector", vector_body.get("file_id", "")) == []
        and len(store.get_chunks_by_file("t1201-vector", survivor_body.get("file_id", ""))) == 1,
        f"deleted={deleted_count}",
    )
    create_kb("t1201-delete")
    def section_file(count):
        return "\n\n".join(f"## Section {index}\nContent {index}" for index in range(count)).encode("utf-8")
    file_a = body_of(upload("t1201-delete", "a.md", section_file(5)))
    file_b = body_of(upload("t1201-delete", "b.md", section_file(3)))
    before_b = chunks_for("t1201-delete", file_b.get("file_id", ""))
    removed = store.delete_by_file("t1201-delete", file_a.get("file_id", ""))
    check("AC-F008-02 literal 5/3 deletion fixture", file_a.get("chunks") == 5 and len(before_b) == 3 and removed == 5 and chunks_for("t1201-delete", file_a.get("file_id", "")) == [] and chunks_for("t1201-delete", file_b.get("file_id", "")) == before_b and store.get_chunk_count("t1201-delete") == 3)

    private_accesses: list[str] = []
    app_root = _BACKEND / "app"
    for source_path in app_root.rglob("*.py"):
        if source_path.name == "vector_store.py":
            continue
        tree = ast.parse(source_path.read_text(encoding="utf-8"), filename=str(source_path))
        for node in ast.walk(tree):
            if isinstance(node, ast.Attribute) and node.attr in {"_client", "_collection"}:
                private_accesses.append(f"{source_path.relative_to(_BACKEND)}:{node.lineno}")
    check(
        "AC-F008-03 application code does not access Chroma private attributes",
        private_accesses == [],
        f"violations={private_accesses}",
    )

    expected_extensions = {
        ".txt",
        ".md",
        ".csv",
        ".json",
        ".log",
        ".pdf",
        ".docx",
        ".xlsx",
        ".xlsm",
        ".xltx",
        ".xltm",
    }
    check(
        "all 11 supported upload extensions complete the ingestion pipeline",
        supported_successes == expected_extensions,
        f"observed={sorted(supported_successes)}",
    )

    failed = [(label, detail) for label, ok, detail in results if not ok]
    print("\n" + "=" * 78)
    print(f"Required checks: {len(results) - len(failed)}/{len(results)} passed")
    print("RESULT:", "PASS" if not failed else "FAIL")
    if failed:
        for label, detail in failed:
            print(f"  - {label}" + (f" — {detail}" if detail else ""))
    print("=" * 78)
    return 1 if failed else 0


def _emit(text: str) -> None:
    """Print child output safely on Windows consoles with legacy codecs."""
    if "--live" in sys.argv:
        from dotenv import dotenv_values
        import os
        values = dotenv_values(_BACKEND / ".env")
        for name in ("DASHSCOPE_API_KEY", "DEEPSEEK_API_KEY"):
            for value in (values.get(name), os.environ.get(name)):
                if value:
                    text = text.replace(value, "[REDACTED]")
    try:
        sys.stdout.write(text)
    except UnicodeEncodeError:
        encoding = sys.stdout.encoding or "utf-8"
        sys.stdout.write(text.encode(encoding, errors="replace").decode(encoding))
    sys.stdout.flush()


def main(*, live: bool = False) -> int:
    """Run the matrix in a child, then remove its isolated temp tree."""
    import os
    probe_root = Path(tempfile.mkdtemp(prefix="t1201_ingestion_"))
    try:
        completed = subprocess.run(
            [sys.executable, "-u", str(Path(__file__).resolve()), "--probe-root", str(probe_root)] + (["--live"] if live else []),
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=900,
            env={**os.environ, "PYTHONIOENCODING": "utf-8"},
        )
        _emit(completed.stdout)
        if completed.stderr.strip() and not live:
            _emit("--- child stderr ---\n" + completed.stderr)
        return_code = completed.returncode
    except subprocess.TimeoutExpired as exc:
        partial = exc.stdout or b""
        _emit(partial.decode("utf-8", errors="replace") if isinstance(partial, bytes) else partial)
        _emit("\nT1201 verification timed out.\n")
        return_code = 1

    for _ in range(10):
        shutil.rmtree(probe_root, ignore_errors=True)
        if not probe_root.exists():
            break
        time.sleep(0.3)
    if probe_root.exists():
        print(f"T1201 cleanup failed: {probe_root}")
        return 1
    print("ISOLATED_STORAGE_CLEANUP: PASS")
    return return_code


if __name__ == "__main__":
    live_mode = "--live" in sys.argv
    if "--probe-root" in sys.argv:
        try:
            raise SystemExit(run_probe(Path(sys.argv[2]), live=live_mode))
        except Exception as exc:
            if not live_mode:
                raise
            print(f"BLOCKED: probe terminated ({type(exc).__name__}); exception text withheld")
            raise SystemExit(1)
    raise SystemExit(main(live=live_mode))
